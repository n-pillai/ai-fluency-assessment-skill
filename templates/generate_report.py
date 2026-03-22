"""
AI Fluency Assessment — .docx Report Generator
Usage: python generate_report.py --date YYYY-MM-DD [--prev-date YYYY-MM-DD]

Reads assessments/YYYY-MM-DD.json and generates AI_Fluency_Assessment_YYYY-MM-DD.docx
in the assessments/ folder. Optionally compares to a previous assessment.

Requires: pip install python-docx
"""

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("python-docx not installed. Run: pip install python-docx")
    sys.exit(1)

# Color palette
NAVY = RGBColor(0x1B, 0x2A, 0x4A)      # #1B2A4A
GOLD = RGBColor(0xC9, 0xA0, 0x2E)       # #C9A02E
LIGHT_GOLD = RGBColor(0xF5, 0xE9, 0xC8) # #F5E9C8
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MID_GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2)

STAGE_LABELS = {
    1: "Stage 1 — Zero or near-zero AI",
    2: "Stage 2 — Coding agent in IDE, permissions on",
    3: "Stage 3 — Agent in IDE, YOLO mode",
    4: "Stage 4 — Wide agent in IDE, code is diffs",
    5: "Stage 5 — CLI, single agent, YOLO",
    6: "Stage 6 — CLI, multi-agent, 3-5 parallel",
    7: "Stage 7 — 10+ agents, hand-managed",
    8: "Stage 8 — Building your own orchestrator",
}


def set_cell_bg(cell, color: RGBColor):
    """Set cell background color."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    hex_color = f"{color.red:02X}{color.green:02X}{color.blue:02X}"
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_paragraph(doc, text="", bold=False, size=10, color=None, align=None, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    return p


def add_heading(doc, text, level=1):
    """Add a styled heading."""
    if level == 1:
        p = add_paragraph(doc, text, bold=True, size=14, color=NAVY, space_before=12, space_after=4)
        # Add bottom border
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), f"{GOLD.red:02X}{GOLD.green:02X}{GOLD.blue:02X}")
        pBdr.append(bottom)
        pPr.append(pBdr)
    elif level == 2:
        add_paragraph(doc, text, bold=True, size=11, color=NAVY, space_before=8, space_after=3)
    elif level == 3:
        add_paragraph(doc, text, bold=True, size=10, color=GOLD, space_before=6, space_after=2)


def add_cover(doc, assessment_date, is_comparison=False):
    """Add cover / header block."""
    # Title block — navy background
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_bg(cell, NAVY)
    cell.width = Inches(6.5)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("AI FLUENCY ASSESSMENT")
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.bold = True
    run.font.color.rgb = WHITE

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(16)
    run2 = p2.add_run(datetime.strptime(assessment_date, "%Y-%m-%d").strftime("%B %-d, %Y") if os.name != "nt" else
                      datetime.strptime(assessment_date, "%Y-%m-%d").strftime("%B %d, %Y").lstrip("0").replace(" 0", " "))
    run2.font.name = "Arial"
    run2.font.size = Pt(13)
    run2.font.color.rgb = LIGHT_GOLD

    doc.add_paragraph()


def add_overview_table(doc, current, previous=None):
    """Add the overview summary table."""
    c4d = current["frameworks"]["anthropic_4d"]
    yegge = current["frameworks"]["yegge"]

    rows = [
        ("Yegge Stage", str(yegge["stage"]), yegge.get("stage_label", STAGE_LABELS.get(int(yegge["stage"]), ""))),
        ("Delegation", c4d["delegation"]["rating"], c4d["delegation"]["delta"]),
        ("Description", c4d["description"]["rating"], c4d["description"]["delta"]),
        ("Discernment", c4d["discernment"]["rating"], c4d["discernment"]["delta"]),
        ("Diligence", c4d["diligence"]["rating"], c4d["diligence"]["delta"]),
    ]

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    # Header row
    headers = ["Dimension", "Rating / Stage", "Delta"]
    for i, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        set_cell_bg(cell, NAVY)
        p = cell.paragraphs[0]
        run = p.add_run(header)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.bold = True
        run.font.color.rgb = WHITE

    for i, (dim, rating, delta) in enumerate(rows):
        row = table.add_row()
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE

        set_cell_bg(row.cells[0], bg)
        r0 = row.cells[0].paragraphs[0].add_run(dim)
        r0.font.name = "Arial"
        r0.font.size = Pt(9)
        r0.bold = True

        set_cell_bg(row.cells[1], bg)
        r1 = row.cells[1].paragraphs[0].add_run(rating)
        r1.font.name = "Arial"
        r1.font.size = Pt(9)

        set_cell_bg(row.cells[2], bg)
        r2 = row.cells[2].paragraphs[0].add_run(delta)
        r2.font.name = "Arial"
        r2.font.size = Pt(9)
        if delta not in ("Baseline", "No change", ""):
            r2.font.color.rgb = GOLD

    doc.add_paragraph()


def add_changes_section(doc, current, previous):
    """Add Changes Since Last Assessment section (only when comparing)."""
    if not previous:
        return

    add_heading(doc, "Changes Since Last Assessment", level=1)

    prev_date = previous["date"]
    curr_date = current["date"]
    days = (datetime.strptime(curr_date, "%Y-%m-%d") - datetime.strptime(prev_date, "%Y-%m-%d")).days
    add_paragraph(doc, f"Assessment interval: {days} days ({prev_date} → {curr_date})", color=MID_GRAY, size=9)

    # Yegge delta
    prev_stage = previous["frameworks"]["yegge"]["stage"]
    curr_stage = current["frameworks"]["yegge"]["stage"]
    add_heading(doc, "Yegge Stage", level=2)
    if curr_stage != prev_stage:
        direction = "Advanced" if curr_stage > prev_stage else "Declined"
        add_paragraph(doc, f"{direction}: Stage {prev_stage} → Stage {curr_stage}", bold=True, color=GOLD)
    else:
        add_paragraph(doc, f"No change: Stage {curr_stage}", color=MID_GRAY)

    # 4D deltas
    add_heading(doc, "4D Framework", level=2)
    for key in ["delegation", "description", "discernment", "diligence"]:
        prev_r = previous["frameworks"]["anthropic_4d"][key]["rating"]
        curr_r = current["frameworks"]["anthropic_4d"][key]["rating"]
        delta = current["frameworks"]["anthropic_4d"][key]["delta"]
        label = key.capitalize()
        if prev_r != curr_r:
            add_paragraph(doc, f"{label}: {prev_r} → {curr_r}", color=GOLD, size=9)
        else:
            add_paragraph(doc, f"{label}: No change ({curr_r})", color=MID_GRAY, size=9)

    # New projects/automation
    new_items = current["inventory"].get("new_since_last", [])
    if new_items:
        add_heading(doc, "New Since Last Assessment", level=2)
        for item in new_items:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(item)
            run.font.name = "Arial"
            run.font.size = Pt(9)

    # Edge changes
    prev_edges = {e["id"]: e for e in previous.get("development_edges", [])}
    curr_edges = {e["id"]: e for e in current.get("development_edges", [])}
    edge_changes = [(curr_edges[eid], prev_edges.get(eid)) for eid in curr_edges
                    if eid in prev_edges and curr_edges[eid]["status"] != prev_edges[eid]["status"]]
    new_edges = [curr_edges[eid] for eid in curr_edges if eid not in prev_edges]

    if edge_changes or new_edges:
        add_heading(doc, "Development Edges", level=2)
        for edge, prev_edge in edge_changes:
            prev_status = prev_edge["status"] if prev_edge else "—"
            add_paragraph(doc, f"Edge {edge['id']}: {edge['title']} — {prev_status} → {edge['status']}",
                          color=GOLD, size=9)
        for edge in new_edges:
            add_paragraph(doc, f"New edge: {edge['title']}", color=NAVY, bold=True, size=9)

    doc.add_paragraph()


def add_4d_section(doc, assessment):
    """Full 4D framework section with evidence."""
    add_heading(doc, "Framework 1 — Anthropic 4D AI Fluency", level=1)
    add_paragraph(doc,
        "Four competencies: Delegation (what to give AI vs. retain), Description (communicating effectively), "
        "Discernment (evaluating outputs critically), Diligence (responsible/ethical use).",
        color=MID_GRAY, size=9, space_after=8)

    competency_descriptions = {
        "delegation": "Deciding which tasks to delegate to AI and calibrating the scope of assistance.",
        "description": "Communicating goals, constraints, format, and context clearly.",
        "discernment": "Evaluating AI outputs for accuracy, completeness, and reasoning quality.",
        "diligence": "Using AI responsibly: data safety, verification, and ethical governance.",
    }

    for key in ["delegation", "description", "discernment", "diligence"]:
        data = assessment["frameworks"]["anthropic_4d"][key]
        add_heading(doc, key.capitalize(), level=2)
        add_paragraph(doc, competency_descriptions[key], color=MID_GRAY, size=9, space_after=2)

        # Rating chip
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f"  {data['rating']}  ")
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.bold = True
        run.font.color.rgb = WHITE
        # Can't do inline background in docx easily — use gold text instead
        run.font.color.rgb = GOLD

        if data.get("delta") and data["delta"] != "Baseline":
            run2 = p.add_run(f"  ({data['delta']})")
            run2.font.name = "Arial"
            run2.font.size = Pt(9)
            run2.font.color.rgb = MID_GRAY

        add_heading(doc, "Evidence", level=3)
        for item in data.get("evidence", []):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(item)
            run.font.name = "Arial"
            run.font.size = Pt(9)

    doc.add_paragraph()


def add_yegge_section(doc, assessment):
    """Yegge stage section with stage table and evidence."""
    add_heading(doc, "Framework 2 — Yegge 8-Stage Developer-Agent Evolution", level=1)
    add_paragraph(doc,
        "Tracks progression from zero AI usage (Stage 1) to building your own agent orchestrator (Stage 8).",
        color=MID_GRAY, size=9, space_after=8)

    yegge = assessment["frameworks"]["yegge"]

    # Stage table
    table = doc.add_table(rows=9, cols=2)
    table.style = "Table Grid"

    # Header
    set_cell_bg(table.cell(0, 0), NAVY)
    set_cell_bg(table.cell(0, 1), NAVY)
    for cell, text in zip(table.rows[0].cells, ["Stage", "Description"]):
        run = cell.paragraphs[0].add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.bold = True
        run.font.color.rgb = WHITE

    current_stage = int(yegge["stage"])
    stage_descs = [
        "Zero or near-zero AI (completions, chat questions)",
        "Coding agent in IDE, permissions on",
        "Agent in IDE, YOLO mode",
        "Wide agent in IDE — code is just diffs",
        "CLI, single agent, YOLO — diffs scroll by",
        "CLI, multi-agent, YOLO — 3-5 parallel instances",
        "10+ agents, hand-managed",
        "Building your own orchestrator",
    ]

    for i, desc in enumerate(stage_descs, 1):
        row = table.rows[i]
        is_current = (i == current_stage)
        bg = LIGHT_GOLD if is_current else (LIGHT_GRAY if i % 2 == 0 else WHITE)
        set_cell_bg(row.cells[0], bg)
        set_cell_bg(row.cells[1], bg)

        r0 = row.cells[0].paragraphs[0].add_run(f"Stage {i}" + (" ◀" if is_current else ""))
        r0.font.name = "Arial"
        r0.font.size = Pt(9)
        r0.bold = is_current
        if is_current:
            r0.font.color.rgb = NAVY

        r1 = row.cells[1].paragraphs[0].add_run(desc)
        r1.font.name = "Arial"
        r1.font.size = Pt(9)
        r1.bold = is_current

    doc.add_paragraph()

    # Current stage details
    add_heading(doc, f"Current Stage: {yegge['stage']} — {yegge.get('stage_label', '')}", level=2)

    if yegge.get("delta") and yegge["delta"] != "Baseline":
        add_paragraph(doc, f"Delta: {yegge['delta']}", color=GOLD, bold=True, size=9)

    add_heading(doc, "Evidence for Current Stage", level=3)
    for item in yegge.get("evidence_for_current", []):
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.font.name = "Arial"
        run.font.size = Pt(9)

    add_heading(doc, "Evidence Pointing Toward Next Stage", level=3)
    for item in yegge.get("evidence_for_next", []):
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.font.name = "Arial"
        run.font.size = Pt(9)

    doc.add_paragraph()


def add_inventory_section(doc, assessment):
    """Project inventory appendix."""
    add_heading(doc, "Appendix — Project & Automation Inventory", level=1)

    inv = assessment["inventory"]

    add_heading(doc, "Claude Code Projects", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Project", "Status", "Autonomy"]):
        set_cell_bg(cell, NAVY)
        run = cell.paragraphs[0].add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.bold = True
        run.font.color.rgb = WHITE

    for i, proj in enumerate(inv.get("claude_code_projects", [])):
        row = table.add_row()
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE
        for cell in row.cells:
            set_cell_bg(cell, bg)
        for cell, val in zip(row.cells, [proj.get("name", ""), proj.get("status", ""), proj.get("autonomy", "")]):
            run = cell.paragraphs[0].add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(9)

    doc.add_paragraph()

    add_heading(doc, "Automated Workflows", level=2)
    table2 = doc.add_table(rows=1, cols=3)
    table2.style = "Table Grid"
    for cell, text in zip(table2.rows[0].cells, ["Workflow", "Trigger", "Description"]):
        set_cell_bg(cell, NAVY)
        run = cell.paragraphs[0].add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.bold = True
        run.font.color.rgb = WHITE

    all_workflows = inv.get("automated_workflows", []) + [
        {"name": t.get("name", ""), "trigger": t.get("trigger", ""), "description": t.get("description", "")}
        for t in inv.get("other_ai_tasks", inv.get("cowork_tasks", []))
    ]
    for i, wf in enumerate(all_workflows):
        row = table2.add_row()
        bg = LIGHT_GRAY if i % 2 == 0 else WHITE
        for cell in row.cells:
            set_cell_bg(cell, bg)
        for cell, val in zip(row.cells, [wf.get("name", ""), wf.get("trigger", ""), wf.get("description", "")]):
            run = cell.paragraphs[0].add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(9)

    doc.add_paragraph()


def add_synthesis_section(doc, assessment):
    """Synthesis and development edges."""
    add_heading(doc, "Synthesis & Development Edges", level=1)
    add_paragraph(doc,
        "Development edges are specific, measurable target behaviors for the next assessment period.",
        color=MID_GRAY, size=9, space_after=8)

    for edge in assessment.get("development_edges", []):
        status_color = {
            "resolved": RGBColor(0x2E, 0x7D, 0x32),
            "in-progress": GOLD,
            "open": NAVY
        }.get(edge.get("status", "open"), NAVY)

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f"Edge {edge['id']}: {edge['title']}")
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.bold = True
        run.font.color.rgb = NAVY

        status_run = p.add_run(f"  [{edge.get('status', 'open').upper()}]")
        status_run.font.name = "Arial"
        status_run.font.size = Pt(9)
        status_run.font.color.rgb = status_color

        add_paragraph(doc, edge.get("description", ""), size=9, color=DARK_GRAY, space_after=2)

        p_target = doc.add_paragraph()
        p_target.paragraph_format.space_after = Pt(8)
        r_label = p_target.add_run("Target behavior: ")
        r_label.font.name = "Arial"
        r_label.font.size = Pt(9)
        r_label.bold = True
        r_label.font.color.rgb = GOLD
        r_body = p_target.add_run(edge.get("target_behavior", ""))
        r_body.font.name = "Arial"
        r_body.font.size = Pt(9)
        r_body.font.color.rgb = DARK_GRAY

    doc.add_paragraph()


def generate_report(assessment_date: str, assessments_dir: Path, previous_date: str = None):
    current_path = assessments_dir / f"{assessment_date}.json"
    if not current_path.exists():
        print(f"Assessment not found: {current_path}")
        sys.exit(1)

    with open(current_path) as f:
        current = json.load(f)

    previous = None
    if previous_date:
        prev_path = assessments_dir / f"{previous_date}.json"
        if prev_path.exists():
            with open(prev_path) as f:
                previous = json.load(f)
        else:
            print(f"Warning: previous assessment not found at {prev_path}")
    else:
        # Auto-find most recent previous
        all_jsons = sorted(assessments_dir.glob("*.json"))
        candidates = [p for p in all_jsons if p.stem < assessment_date]
        if candidates:
            with open(candidates[-1]) as f:
                previous = json.load(f)
            print(f"Comparing to previous assessment: {candidates[-1].stem}")

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Set default font
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)

    add_cover(doc, assessment_date, is_comparison=bool(previous))
    add_overview_table(doc, current, previous)

    if previous:
        add_changes_section(doc, current, previous)

    add_heading(doc, "Data Sources", level=1)
    add_paragraph(doc, "This assessment draws from three data channels:", size=9, color=MID_GRAY)
    sources = [
        "Claude Code self-assessment: projects, autonomy levels, automation inventory, Yegge estimate",
        "Other AI tools: automations, integrations, and AI-generated artifacts",
        "Claude.ai manual notes: delegation patterns, error-catching, format specifications",
    ]
    for s in sources:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(s)
        run.font.name = "Arial"
        run.font.size = Pt(9)
    doc.add_paragraph()

    add_4d_section(doc, current)
    add_yegge_section(doc, current)
    add_synthesis_section(doc, current)
    add_inventory_section(doc, current)

    output_path = assessments_dir / f"AI_Fluency_Assessment_{assessment_date}.docx"
    doc.save(str(output_path))
    print(f"Report saved: {output_path}")
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Generate AI Fluency Assessment .docx report")
    parser.add_argument("--date", required=True, help="Assessment date (YYYY-MM-DD)")
    parser.add_argument("--prev-date", help="Previous assessment date for comparison (YYYY-MM-DD)")
    args = parser.parse_args()

    # Locate assessments dir relative to this script
    script_dir = Path(__file__).parent
    assessments_dir = script_dir.parent / "assessments"

    if not assessments_dir.exists():
        print(f"Assessments directory not found: {assessments_dir}")
        sys.exit(1)

    generate_report(args.date, assessments_dir, args.prev_date)


if __name__ == "__main__":
    main()
